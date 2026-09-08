# scripts/05_export_excel.py
"""
Step 5: Excel workbook.

Writes outputs/scotland_route_analysis.xlsx with three sheets:
  Town Scores     all 174 shortlisted towns, their inputs and scores
  Route Schedule  the 14-day route, one town per day
  Summary         intentionally near-empty, for manual pivot tables and charts

Formatting follows CLAUDE.md section 6: a single accent colour on a bold frozen
header row, autofilter, explicit column widths, correct number formats, a colour
scale on the Score column only, a named range over each data table, and A4
landscape print setup with a repeating header row. No openpyxl Table objects, so
the file behaves consistently in LibreOffice Calc.
"""

from pathlib import Path

import geopandas as gpd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

PROCESSED_DIR = Path("data/processed")
OUT_XLSX = Path("outputs/scotland_route_analysis.xlsx")

SCORED_GPKG = PROCESSED_DIR / "towns_scored.gpkg"
ROUTE_GPKG = PROCESSED_DIR / "route.gpkg"

# One accent colour, used on header rows. ARGB with an explicit opaque alpha so
# the fill renders solid in LibreOffice as well as Excel.
ACCENT = "FF1F4E79"
SCORE_SCALE_LOW = "FFFFFFFF"
SCORE_SCALE_HIGH = "FF9DC3E6"
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)


def style_data_sheet(ws, n_cols, n_rows, widths, number_formats):
    """Header styling, widths, number formats, freeze, autofilter and print setup
    shared by the two data sheets. number_formats maps 1-based column index to a
    format string."""
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for col, fmt in number_formats.items():
        for row in range(2, n_rows + 2):
            ws.cell(row=row, column=col).number_format = fmt

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def build_town_scores(ws, towns):
    headers = [
        "Town", "Council area", "Population", "Share aged 55+",
        "Distance to nearest competitor (km)", "Population score",
        "Age 55+ score", "Competitor distance score", "Score",
    ]
    ws.append(headers)
    for _, t in towns.iterrows():
        ws.append([
            t["name"], t["council_area"], int(t["population"]),
            float(t["share_55_plus"]), t["comp_dist_m"] / 1000,
            float(t["norm_population"]), float(t["norm_share_55_plus"]),
            float(t["norm_comp_dist"]), float(t["score"]),
        ])

    n_rows = len(towns)
    style_data_sheet(
        ws, n_cols=len(headers), n_rows=n_rows,
        widths=[24, 22, 12, 14, 20, 15, 13, 18, 9],
        number_formats={3: "#,##0", 4: "0.0%", 5: "0.0",
                        6: "0.000", 7: "0.000", 8: "0.000", 9: "0.000"},
    )
    # Colour scale on the Score column only (column I).
    ws.conditional_formatting.add(
        f"I2:I{n_rows + 1}",
        ColorScaleRule(start_type="min", start_color=SCORE_SCALE_LOW,
                       end_type="max", end_color=SCORE_SCALE_HIGH),
    )
    return n_rows


def build_route_schedule(ws, stops, loop_km, loop_min):
    headers = [
        "Day", "Town", "Council area", "Population",
        "Distance from previous (km)", "Driving time from previous (min)",
        "Cumulative distance (km)",
    ]
    ws.append(headers)
    for _, s in stops.iterrows():
        first_day = s["day"] == 1
        ws.append([
            int(s["day"]), s["name"], s["council_area"], int(s["population"]),
            None if first_day else round(float(s["leg_km"]), 1),
            None if first_day else round(float(s["leg_min"])),
            round(float(s["cum_km"]), 1),
        ])

    n_rows = len(stops)
    style_data_sheet(
        ws, n_cols=len(headers), n_rows=n_rows,
        widths=[6, 24, 22, 12, 22, 26, 22],
        number_formats={1: "0", 4: "#,##0", 5: "0.0", 6: "0", 7: "0.0"},
    )
    # Loop total, two rows below the table so it stays outside the named range.
    total_row = n_rows + 3
    ws.cell(row=total_row, column=2, value="Loop total (including return to start)").font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=round(loop_km, 1)).number_format = "0.0"
    ws.cell(row=total_row, column=6, value=round(loop_min)).number_format = "0"
    ws.cell(row=total_row + 1, column=2, value=f"= {loop_min / 60:.1f} hours driving")
    return n_rows


def build_summary(ws, n_towns, n_stops, loop_km, loop_min, entry_town):
    ws.column_dimensions["A"].width = 100
    ws["A1"] = "Summary"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "This sheet is intentionally left mostly blank. Build pivot tables and charts here,",
        "using the named ranges defined on the other two sheets:",
        "",
        f"    TownScores      all {n_towns} shortlisted towns and their scores  (sheet 'Town Scores')",
        f"    RouteSchedule   the {n_stops}-day route                            (sheet 'Route Schedule')",
        "",
        "Route: closed loop, one town per day, "
        f"{loop_km:,.0f} km / {loop_min / 60:.1f} h driving. Entry point {entry_town}.",
        "Scores combine population (log scale), share aged 55+, and distance to the nearest",
        "competitor, each min-max normalised to 0-1 with equal weight.",
    ]
    for i, text in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=text)


def main():
    towns = gpd.read_file(SCORED_GPKG, layer="towns_scored").sort_values(
        "score", ascending=False
    )
    stops = gpd.read_file(ROUTE_GPKG, layer="route_stops").sort_values("day")
    line = gpd.read_file(ROUTE_GPKG, layer="route_line").iloc[0]
    loop_km = float(line["distance_km"])
    loop_min = float(line["duration_min"])
    entry_town = stops.iloc[0]["name"]

    wb = Workbook()
    ws_scores = wb.active
    ws_scores.title = "Town Scores"
    ws_schedule = wb.create_sheet("Route Schedule")
    ws_summary = wb.create_sheet("Summary")

    n_scores = build_town_scores(ws_scores, towns)
    n_schedule = build_route_schedule(ws_schedule, stops, loop_km, loop_min)
    build_summary(ws_summary, len(towns), len(stops), loop_km, loop_min, entry_town)

    wb.defined_names["TownScores"] = DefinedName(
        "TownScores", attr_text=f"'Town Scores'!$A$1:$I${n_scores + 1}"
    )
    wb.defined_names["RouteSchedule"] = DefinedName(
        "RouteSchedule", attr_text=f"'Route Schedule'!$A$1:$G${n_schedule + 1}"
    )

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)

    print(f"Wrote {OUT_XLSX}")
    print(f"  Town Scores    {n_scores} towns")
    print(f"  Route Schedule {n_schedule} days, loop total {loop_km:,.0f} km / {loop_min / 60:.1f} h")
    print(f"  Summary        note plus named ranges TownScores, RouteSchedule")


if __name__ == "__main__":
    main()
